"""Build a source-attributed contextual vocabulary collection, without model calls.

Source archives stay in ignored app/data/lexicon-sources. Nothing in this script
reads or writes learner progress, book lesson caches, or runtime processes.
"""
from __future__ import annotations

import argparse
import bz2
import csv
import gzip
import hashlib
import io
import json
import re
import urllib.request
from urllib.parse import quote
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
CACHE = APP / "data" / "lexicon-sources"
OUT = APP / "content" / "lexicon"
VERSION = "context-lexicon-v1"
NGSL_PAGE = "https://www.newgeneralservicelist.com/new-general-service-list"
TATOEBA_PAGE = "https://tatoeba.org/en/downloads"
KAIIKI_PAGE = "https://kaikki.org/dictionary/rawdata.html"
SOURCES = [
    {"id": "ngsl-31k", "title": "NGSL plus statistics and frequencies to 31k", "file": "NGSLwithSFI-31K.xlsx", "url": "https://www.newgeneralservicelist.com/s/NGSLwithSFI-31K.xlsx", "page": NGSL_PAGE, "license": "CC-BY-SA-4.0", "author": "Charles Browne, Brent Culligan, Joseph Phillips", "maxBytes": 6000000},
    {"id": "ngsl-1.2", "title": "NGSL 1.2 with basic statistics", "file": "NGSL_12_stats.csv", "url": "https://www.newgeneralservicelist.com/s/NGSL_12_stats.csv", "page": NGSL_PAGE, "license": "CC-BY-SA-4.0", "author": "Charles Browne, Brent Culligan, Joseph Phillips", "maxBytes": 1000000},
    {"id": "ngsl-gr-1.0", "title": "NGSL-Graded Reader 1.0", "file": "NGSL-GR_rank.csv", "url": "https://www.newgeneralservicelist.com/s/NGSL-GR_rank.csv", "page": "https://www.newgeneralservicelist.com/ngsl-graded-reader", "license": "CC-BY-SA-4.0", "author": "Charles Browne, Brent Culligan", "maxBytes": 1000000},
    {"id": "ngsl-spoken-1.2", "title": "NGSL-Spoken 1.2 with statistics", "file": "NGSL-Spoken_12_stats.csv", "url": "https://www.newgeneralservicelist.com/s/NGSL-Spoken_12_stats.csv", "page": "https://www.newgeneralservicelist.com/ngsl-spoken", "license": "CC-BY-SA-4.0", "author": "Charles Browne, Brent Culligan", "maxBytes": 1000000},
    {"id": "ngsl-supplement", "title": "NGSL 52 supplementary words", "file": "SUP_lemmatized.csv", "url": "https://www.newgeneralservicelist.com/s/SUP_lemmatized.csv", "page": NGSL_PAGE, "license": "CC-BY-SA-4.0", "author": "Charles Browne, Brent Culligan, Joseph Phillips", "maxBytes": 1000000},
    {"id": "nawl-1.2", "title": "NAWL 1.2 alphabetized", "file": "NAWL_12_alphabetized_description.txt", "url": "https://www.newgeneralservicelist.com/s/NAWL_12_alphabetized_description.txt", "page": "https://www.newgeneralservicelist.com/new-academic-word-list", "license": "CC-BY-SA-4.0", "author": "Charles Browne, Brent Culligan, Joseph Phillips", "maxBytes": 1000000},
    {"id": "tatoeba-eng", "title": "Tatoeba English detailed sentences", "file": "eng_sentences_detailed.tsv.bz2", "url": "https://downloads.tatoeba.org/exports/per_language/eng/eng_sentences_detailed.tsv.bz2", "page": TATOEBA_PAGE, "license": "CC-BY-2.0-FR", "author": "Tatoeba contributors; individual authors retained per sentence", "maxBytes": 55000000},
    {"id": "tatoeba-rus", "title": "Tatoeba Russian detailed sentences", "file": "rus_sentences_detailed.tsv.bz2", "url": "https://downloads.tatoeba.org/exports/per_language/rus/rus_sentences_detailed.tsv.bz2", "page": TATOEBA_PAGE, "license": "CC-BY-2.0-FR", "author": "Tatoeba contributors; individual authors retained per sentence", "maxBytes": 40000000},
    {"id": "tatoeba-eng-rus-links", "title": "Tatoeba direct English-Russian translation links", "file": "eng-rus_links.tsv.bz2", "url": "https://downloads.tatoeba.org/exports/per_language/eng/eng-rus_links.tsv.bz2", "page": TATOEBA_PAGE, "license": "CC-BY-2.0-FR", "author": "Tatoeba contributors", "maxBytes": 10000000},
    {"id": "kaikki-simple", "title": "Kaikki raw Simple English Wiktionary extraction", "file": "simple-extract.jsonl.gz", "url": "https://kaikki.org/dictionary/downloads/simple/simple-extract.jsonl.gz", "page": KAIIKI_PAGE, "license": "CC-BY-SA-4.0", "author": "Simple English Wiktionary contributors; extraction by Wiktextract/Kaikki", "maxBytes": 12000000},
]


def atomic_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def file_sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def download_source(source: dict, refresh: bool = False) -> dict:
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / source["file"]
    receipt_path = CACHE / (source["file"] + ".receipt.json")
    if path.exists() and receipt_path.exists() and not refresh:
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        if file_sha(path) != receipt["sha256"]:
            raise ValueError(f"Cached source hash mismatch: {path.name}")
        return receipt
    req = urllib.request.Request(source["url"], headers={"User-Agent": "EnglishLearningContextLexicon/1.0 (source attribution retained)"})
    tmp = path.with_suffix(path.suffix + ".part")
    try:
        with urllib.request.urlopen(req, timeout=45) as response, tmp.open("wb") as target:
            content_length = response.headers.get("Content-Length")
            if content_length and int(content_length) > source["maxBytes"]:
                raise ValueError(f"Source exceeds bounded download: {source['id']}")
            count = 0
            for block in iter(lambda: response.read(1024 * 1024), b""):
                count += len(block)
                if count > source["maxBytes"]:
                    raise ValueError(f"Source exceeds bounded download: {source['id']}")
                target.write(block)
            receipt = {**source, "retrievedAt": datetime.now(timezone.utc).isoformat(), "resolvedUrl": response.url, "lastModified": response.headers.get("Last-Modified"), "bytes": count}
        tmp.replace(path)
        receipt["sha256"] = file_sha(path)
        atomic_json(receipt_path, receipt)
        print(f"Downloaded {source['id']}: {count:,} bytes", flush=True)
        return receipt
    finally:
        if tmp.exists():
            tmp.unlink()


def download_sources(refresh=False) -> list[dict]:
    with ThreadPoolExecutor(max_workers=4) as pool:
        receipts = list(pool.map(lambda source: download_source(source, refresh), SOURCES))
    # Downloading is not publication: keep the live attribution registry intact
    # until a complete validated build has included its editorial sources too.
    atomic_json(CACHE / "download-receipts.json", {"version": VERSION, "sources": receipts})
    return receipts


def normalized_word(value: str) -> str:
    return value.strip().replace("’", "'").casefold()


def load_ranked_words() -> list[dict]:
    """Keep actual column semantics; neither worksheet order nor WL rank is CEFR."""
    from openpyxl import load_workbook  # Read only; source workbook is never saved.
    cached = CACHE / "ranked-words.json"
    sha = file_sha(CACHE / "NGSLwithSFI-31K.xlsx")
    if cached.exists():
        data = json.loads(cached.read_text(encoding="utf-8"))
        if data.get("sourceSHA256") == sha:
            return data["words"]
    workbook = load_workbook(CACHE / "NGSLwithSFI-31K.xlsx", read_only=True, data_only=True)
    words = {}
    for source_row, row in enumerate(workbook.active.iter_rows(min_row=2, values_only=True), 2):
        lemma, _, _, sfi, _, _, _, raw_rank = row[:8]
        if not isinstance(lemma, str) or not isinstance(raw_rank, (int, float)):
            continue
        key = normalized_word(lemma)
        if not re.fullmatch(r"[a-z]+(?:[-'][a-z]+)*", key):
            continue
        item = {"key": key, "word": lemma.strip(), "rank": {"sourceId": "ngsl-31k", "value": int(raw_rank), "metric": "RawFreq_Rank", "sourceRow": source_row}, "sfi": sfi}
        if key not in words or item["rank"]["value"] < words[key]["rank"]["value"]:
            words[key] = item
    workbook.close()
    result = sorted(words.values(), key=lambda x: (x["rank"]["value"], x["key"]))
    atomic_json(cached, {"sourceSHA256": sha, "words": result})
    return result


def memberships() -> dict[str, list[dict]]:
    result = defaultdict(list)
    for file, source_id, word_column, rank_column, metric in [
        ("NGSL_12_stats.csv", "ngsl-1.2", "Lemma", "SFI Rank", "SFI Rank"),
        ("NGSL-GR_rank.csv", "ngsl-gr-1.0", "Word", "WordID", "WordID (graded-reader sequence, not CEFR)"),
        ("NGSL-Spoken_12_stats.csv", "ngsl-spoken-1.2", "Lemma", "Rank", "Rank"),
    ]:
        with (CACHE / file).open(encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                raw_rank = row.get(rank_column, "")
                value = int(raw_rank) if raw_rank.isdigit() else None
                result[normalized_word(row[word_column])].append({"sourceId": source_id, "value": value, "metric": metric, **({"sourceRankValue": raw_rank} if value is None else {})})
    with (CACHE / "SUP_lemmatized.csv").open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            if row and row[0].strip():
                result[normalized_word(row[0])].append({"sourceId": "ngsl-supplement", "value": None, "metric": "membership"})
    for line in (CACHE / "NAWL_12_alphabetized_description.txt").read_text(encoding="utf-8-sig").splitlines():
        if re.fullmatch(r"[a-z]+(?:[-'][a-z]+)*", line.strip()):
            result[normalized_word(line)].append({"sourceId": "nawl-1.2", "value": None, "metric": "membership"})
    return result


def lexical_id(word: str) -> str:
    # A hash suffix avoids collisions after punctuation stripping.
    key = normalized_word(word)
    slug = re.sub(r"[^a-z0-9]+", "-", key).strip("-")
    return "lex-" + slug + "-" + hashlib.sha256(key.encode()).hexdigest()[:8]


def target_spans(text: str, word: str, forms: list[str] | None = None) -> list[dict]:
    """Offsets use UTF-16 code units, matching JavaScript String.slice."""
    keys = {normalized_word(word), *(normalized_word(form) for form in forms or [])}
    spans = []
    for match in re.finditer(r"[^\W\d_]+(?:['’\-][^\W\d_]+)*", text):
        if normalized_word(match.group()) in keys:
            start = len(text[:match.start()].encode("utf-16-le")) // 2
            end = len(text[:match.end()].encode("utf-16-le")) // 2
            spans.append({"start": start, "end": end, "text": match.group()})
    return spans


def load_dictionary() -> dict[str, list[dict]]:
    result = defaultdict(list)
    with gzip.open(CACHE / "simple-extract.jsonl.gz", "rt", encoding="utf-8") as handle:
        for line in handle:
            item = json.loads(line)
            if item.get("lang_code") != "en" or item.get("pos") in {"name", "hard-redirect", "symbol", "prefix", "suffix"}:
                continue
            word = item.get("word", "")
            key = normalized_word(word)
            url = "https://simple.wiktionary.org/wiki/" + quote(word.replace(" ", "_"), safe="")
            for index, sense in enumerate(item.get("senses", [])):
                glosses = sense.get("glosses", [])
                if not glosses or sense.get("form_of"):
                    continue
                definition = " ".join(glosses)
                if re.match(r"^(?:The )?(?:plural|past|present|third.person|comparative|superlative|alternative (?:spelling|form))\b", definition, re.I):
                    continue
                source = {"sourceId": "kaikki-simple", "url": url, "license": "CC-BY-SA-4.0", "attribution": "Simple English Wiktionary contributors", "sourceSenseIndex": index}
                identity = hashlib.sha256((item['pos'] + '\0' + definition).encode()).hexdigest()[:12]
                examples = []
                forms = [form['form'] for form in item.get('forms', []) if isinstance(form.get('form'), str) and re.fullmatch(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", form['form']) and not any(tag in form.get('tags', []) for tag in ['alternative', 'archaic', 'obsolete'])]
                for example in sense.get("examples", []):
                    # Externally quoted material has its own rights, not a blanket CC licence.
                    if example.get("ref") or example.get("type") == "quotation":
                        continue
                    text = example.get("text", "")
                    if target_spans(text, word, forms) and 5 <= len(text.split()) <= 35:
                        examples.append({"en": text, "targetSpans": target_spans(text, word, forms)})
                pronunciations = []
                for sound in item.get("sounds", []):
                    if sound.get("ipa"):
                        pronunciations.append({"ipa": sound["ipa"], "regionTags": sound.get("tags", []), "rawTags": sound.get("raw_tags", []), "source": source})
                result[key].append({"id": lexical_id(word) + "-sense-" + identity, "pos": item["pos"], "definition": definition, "tags": sense.get("tags", []), "rawTags": sense.get("raw_tags", []), "categories": sense.get("categories", []), "pronunciations": pronunciations, "examples": examples, "source": source, "quality": "source-imported; not editor-verified"})
    # Several source pages/capitalizations can repeat the same POS and gloss.
    # Preserve their provenance while exposing one stable semantic ID.
    for word, senses in result.items():
        unique = {}
        for sense in senses:
            if sense['id'] not in unique:
                unique[sense['id']] = sense
                continue
            kept = unique[sense['id']]
            occurrences = kept.setdefault('sourceOccurrences', [kept['source']])
            if sense['source'] not in occurrences:
                occurrences.append(sense['source'])
            for field in ['examples', 'pronunciations', 'tags', 'rawTags', 'categories']:
                for value in sense[field]:
                    if value not in kept[field]:
                        kept[field].append(value)
        result[word] = list(unique.values())
    return result


def add_dictionary_contexts(contexts: dict, dictionary: dict) -> dict:
    """Only examples nested under an actual source sense; never guess Tatoeba senses."""
    result = {word: list(values) for word, values in contexts.items()}
    for word, senses in dictionary.items():
        if result.get(word):
            continue
        options = []
        for sense in senses:
            for example in sense.get('examples', []):
                text = example['en']
                if not text or len(text) > 350 or any(s in text for s in ['http', '<', '>', '\n']):
                    continue
                # A dictionary example can legitimately be shorter than a standalone
                # sentence-corpus example: its sense is supplied by the source.
                if re.search(r"\b(?:fuck\w*|shit\w*|nigg\w*|fagg\w*|porn\w*)\b", text, re.I):
                    continue
                context_id = 'wiktionary-' + hashlib.sha256((sense['id'] + '\0' + text).encode()).hexdigest()[:16]
                options.append({'id': context_id, 'en': text, 'ru': None, 'targetSpans': example['targetSpans'], 'senseId': sense['id'], 'variety': 'unverified', 'source': sense['source'], 'alignment': 'example nested under this sense in source dictionary; not editor-verified', 'quality': 'source-linked-context'})
        if options:
            # Context order is deterministic; it has no pedagogic/CEFR significance.
            result[word] = options[:2]
    return result


def editorial_documents() -> list[tuple[Path, dict]]:
    paths = [OUT / 'editorial-contexts.json', *sorted(OUT.glob('editorial-core-*.json'))]
    return [(path, data) for path in paths if (data := json.loads(path.read_text(encoding='utf-8'))).get('approved', True)]


def add_editorial_contexts(contexts: dict, dictionary: dict) -> dict:
    result = contexts
    for _, data in editorial_documents():
        result = add_editorial_document(result, dictionary, data)
    return result


def add_editorial_document(contexts: dict, dictionary: dict, data: dict) -> dict:
    metadata = data['source']
    result = {word: list(values) for word, values in contexts.items()}
    for row in data['entries']:
        word = normalized_word(row['word'])
        source = {'sourceId': metadata['id'], 'license': metadata['license'], 'attribution': metadata['author'], 'checkedAt': metadata['checkedAt'], 'kind': 'original-editorial'}
        sense_id = lexical_id(word) + '-sense-editorial-1'
        context_id = lexical_id(word) + '-context-editorial-1'
        sense = {'id': sense_id, 'pos': row['pos'], 'definition': row['meaningEn'], 'definitionRu': row['meaningRu'], 'tags': row.get('registerTags', []), 'rawTags': [], 'categories': [], 'pronunciations': [], 'examples': [], 'source': source, 'quality': 'editor-reviewed-for-this-context', 'verificationSources': row.get('verificationSources', [])}
        sense['lexicalType'] = row.get('lexicalType', 'word')
        sense['displayHeadword'] = row.get('displayHeadword', row['word'])
        dictionary[word] = [sense, *dictionary.get(word, [])]
        if word in {'shareholder', 'fascinate'}:
            issue = {'shareholder': 'Source says shares stocks; the intended ownership sense requires owns shares.', 'fascinate': 'Source describes a person or thing rather than the meaning of the verb.'}[word]
            for original in dictionary[word][1:]:
                original['sourceIssues'] = [issue]
                original['quality'] = 'source-issue; consult the editorial sense for this context'
        if word == 'stair':
            for original in dictionary[word][1:]:
                original['sourceNotes'] = [{'note': 'The source selects the flight-of-steps sense. The editorial bottom stair example uses the separate single-step sense; both uses are attested.', 'verificationUrl': 'https://www.merriam-webster.com/dictionary/stair', 'checkedAt': '2026-09-10'}]
        context = {'id': context_id, 'en': row['en'], 'ru': row['ru'], 'targetSpans': target_spans(row['en'], row['word'], row.get('targetForms')), 'senseId': sense_id, 'variety': row.get('variety', 'en-US-compatible'), 'source': source, 'translationSource': source, 'alignment': 'original example and Russian translation reviewed together', 'quality': 'context-reviewed', 'explanation': row['explanation'], 'productionTask': row['task'], 'collocations': [{'text': row['phrase'], 'evidence': 'used in this original example; no corpus-frequency claim', 'contextId': context_id}], 'topics': [row['topic']], 'registerTags': row.get('registerTags', []), 'usAlternative': row.get('usAlternative')}
        if not context['targetSpans']:
            raise ValueError(f'Editorial example lacks its target word: {word}')
        if row['phrase'].casefold() not in row['en'].casefold():
            raise ValueError(f'Editorial phrase is absent from example: {word}')
        if row.get('targetForms'):
            context['targetForms'] = row['targetForms']
            context['targetRelation'] = 'explicit editorial surface form; see the explanation for inflection, spelling or word formation'
        result[word] = [context, *result.get(word, [])]
    return result


def sentence_score(text: str) -> int | None:
    """A transparent shortlist filter, not semantic validation of a sentence."""
    words = re.findall(r"[A-Za-z]+(?:['’\-][A-Za-z]+)*", text)
    if not 7 <= len(words) <= 28 or len(text) > 240:
        return None
    if not text or not text[0].isupper() or text[-1] not in ".?!":
        return None
    if any(marker in text for marker in ["http", "www.", "@", "{", "}", "<", ">", "\n", "\t"]):
        return None
    if re.search(r"\b(?:this|that|it) is (?:a |an |the )?[A-Za-z'-]+[.!?]$", text, re.I) and len(words) < 9:
        return None
    if re.search(r"\b(?:fuck\w*|shit\w*|nigg\w*|fagg\w*|porn\w*)\b", text, re.I):
        return None  # Sensitive/register-marked material needs its own editorial selection.
    if len(set(w.casefold() for w in words)) < len(words) * 0.58:
        return None
    score = 100 - abs(len(words) - 13) * 2
    score -= 5 * len(re.findall(r"\b(?:Tom|Mary|John|Boston)\b", text))
    score -= 10 * len(re.findall(r"[\"\[\];]", text))
    return score


def iter_tsv(filename):
    with bz2.open(CACHE / filename, "rt", encoding="utf-8", newline="") as handle:
        # Sentence text is literal TSV, not CSV with quote escaping.
        for line in handle:
            yield line.rstrip("\r\n").split("\t")


def sentence_source(source_id, row):
    return {"sourceId": source_id, "sentenceId": int(row[0]), "author": row[3] if row[3] != "\\N" else None, "url": "https://tatoeba.org/en/sentences/show/" + row[0], "license": "CC-BY-2.0-FR", "modifiedAt": row[5]}


def collect_contexts(words: set[str], shortlist_size=4) -> tuple[dict, dict]:
    translated = {int(row[0]) for row in iter_tsv("eng-rus_links.tsv.bz2")}
    best = defaultdict(list)
    checked = eligible = 0
    for row in iter_tsv("eng_sentences_detailed.tsv.bz2"):
        checked += 1
        if len(row) != 6 or row[1] != "eng" or int(row[0]) not in translated:
            continue
        score = sentence_score(row[2])
        if score is None:
            continue
        eligible += 1
        matched = set(normalized_word(m.group()) for m in re.finditer(r"[^\W\d_]+(?:['’\-][^\W\d_]+)*", row[2])) & words
        for word in matched:
            candidate = (score, -int(row[0]), row)
            selected = best[word]
            if len(selected) < shortlist_size or candidate[:2] > selected[-1][:2]:
                selected.append(candidate)
                selected.sort(key=lambda x: x[:2], reverse=True)
                del selected[shortlist_size:]
    del translated
    ids = {int(candidate[2][0]) for candidates in best.values() for candidate in candidates}
    links = defaultdict(list)
    for en_id, ru_id in iter_tsv("eng-rus_links.tsv.bz2"):
        if int(en_id) in ids:
            links[int(en_id)].append(int(ru_id))
    needed_ru = {r for rs in links.values() for r in rs}
    ru_rows = {}
    for row in iter_tsv("rus_sentences_detailed.tsv.bz2"):
        if len(row) == 6 and int(row[0]) in needed_ru and re.search("[А-Яа-яЁё]", row[2]):
            if len(row[2]) < 400 and not any(s in row[2] for s in ["http", "www.", "\t", "\n"]):
                ru_rows[int(row[0])] = row
    contexts = {}
    for word, candidates in best.items():
        values = []
        used = set()
        for score, _, row in candidates:
            translations = [ru_rows[r] for r in sorted(links[int(row[0])]) if r in ru_rows]
            if not translations or row[2] in used:
                continue
            translated_row = translations[0]
            used.add(row[2])
            values.append({"id": "tatoeba-" + row[0] + "-" + translated_row[0], "en": row[2], "ru": translated_row[2], "targetSpans": target_spans(row[2], word), "senseId": None, "variety": "unverified", "source": sentence_source("tatoeba-eng", row), "translationSource": sentence_source("tatoeba-rus", translated_row), "alignment": "direct source translation link; not editor-verified", "quality": "imported-context", "selectionScore": score})
        if values:
            contexts[word] = values[:2]
    return contexts, {"englishRowsScanned": checked, "eligibleLinkedSentences": eligible, "wordsWithContext": len(contexts)}


def validate_entries(document: dict, sources: list[dict]) -> None:
    if document.get('spanEncoding') != 'utf-16':
        raise ValueError('Unknown target-span encoding')
    known_sources = {s['id'] for s in sources}
    seen_ids, seen_words = set(), set()
    for entry in document['entries']:
        key = normalized_word(entry['word'])
        if entry['id'] in seen_ids or key in seen_words:
            raise ValueError('Duplicate word or entry ID')
        seen_ids.add(entry['id']); seen_words.add(key)
        if entry.get('cefr') is not None:
            raise ValueError('This import contains no verified CEFR source mapping')
        senses = {s['id']: s for s in entry.get('senses', [])}
        if len(senses) != len(entry.get('senses', [])):
            raise ValueError(f'Duplicate sense ID: {key}')
        if not entry.get('contexts'):
            raise ValueError(f'Entry has no context: {key}')
        for context in entry['contexts']:
            source_id = context['source']['sourceId']
            if source_id not in known_sources:
                raise ValueError(f'Unknown context source: {source_id}')
            if context.get('senseId') and context['senseId'] not in senses:
                raise ValueError(f'Context points to missing sense: {key}')
            if document.get('unicodeTokenSelectionValidated') and source_id == 'tatoeba-eng' and not context.get('excludedFromStudy') and context['targetSpans'] != target_spans(context['en'], entry['word']):
                raise ValueError(f'Tatoeba target must be a complete Unicode token: {key}')
            if context.get('excludedFromStudy') and (not context.get('exclusionReason') or context.get('selectionReviewSourceId') not in known_sources):
                raise ValueError(f'Excluded source context needs an attributed reason: {key}')
            if context.get('quality') == 'ai-context-reviewed':
                review = context.get('completionReview', {})
                selected = senses.get(context.get('senseId'), {})
                if (review.get('sourceId') not in known_sources or
                        selected.get('source', {}).get('sourceId') != review.get('sourceId') or
                        not isinstance(review.get('reviewPasses'), int) or review['reviewPasses'] < 1 or
                        not re.fullmatch('[a-f0-9]{64}', review.get('baseSHA256', ''))):
                    raise ValueError('AI context alignment requires a selected editorial sense and review receipt')
            if source_id == 'tatoeba-eng' and context.get('senseId') and context['quality'] not in {'context-reviewed', 'ai-context-reviewed'}:
                raise ValueError('Tatoeba sense alignment must not be inferred automatically')
            if context.get('ru') and context.get('translationSource', {}).get('sourceId') not in known_sources:
                raise ValueError(f'Russian context has no attribution: {key}')
            raw = context['en'].encode('utf-16-le')
            if not context.get('targetSpans'):
                raise ValueError(f'Context has no target highlight: {key}')
            for span in context['targetSpans']:
                if not 0 <= span['start'] < span['end'] <= len(raw) // 2:
                    raise ValueError(f'Invalid target-span bounds: {key}')
                if raw[span['start'] * 2:span['end'] * 2].decode('utf-16-le') != span['text']:
                    raise ValueError(f'Target-span text mismatch: {key}')
        for sense in senses.values():
            if sense['source']['sourceId'] not in known_sources:
                raise ValueError(f'Unknown dictionary source: {key}')


def select_words(eligible, member_index, limit, previous_words=()):
    available = {item['key']: item for item in eligible}
    missing_previous = set(previous_words) - set(available)
    if missing_previous:
        raise ValueError('Source refresh would remove published contexts; review before replacing: ' + ', '.join(sorted(missing_previous)[:20]))
    chosen = {key: available[key] for key in previous_words}
    # Required vocabulary adds to the published bank, never displaces saved IDs.
    for item in eligible:
        if item['key'] in member_index:
            chosen[item['key']] = item
    for item in eligible:
        if len(chosen) >= limit:
            break
        chosen.setdefault(item['key'], item)
    return sorted(chosen.values(), key=lambda w: (w['rank']['value'] if w['rank']['value'] is not None else 999999, w['key']))


def retain_published_contexts(contexts, dictionary, previous_entries):
    """Keep stable source records when an improved selector stops choosing them.

    Repairs decide whether an old record remains usable; selection alone cannot
    delete a context referenced by learner history or add unreviewed fallback
    examples to a completed entry. Explicit overlays add replacement contexts. A
    refreshed source with the same ID wins, so overlay hashes still detect an
    actual source edit rather than silently replacing it with the old quotation.
    """
    for entry in previous_entries:
        key = normalized_word(entry['word'])
        selected = {context['id']: context for context in contexts.get(key, [])}
        current = [selected.get(context['id'], json.loads(json.dumps(context))) for context in entry['contexts']]
        contexts[key] = current
        required_senses = {context.get('senseId') for context in current}
        known_senses = {sense['id'] for sense in dictionary.get(key, [])}
        for sense in entry.get('senses', []):
            if sense['id'] in required_senses and sense['id'] not in known_senses:
                dictionary.setdefault(key, []).append(json.loads(json.dumps(sense)))
                known_senses.add(sense['id'])


def build(limit=10000, pilot=False):
    words = load_ranked_words()
    dictionary = load_dictionary()
    member_index = memberships()
    cache = CACHE / "context-candidates.json"
    fingerprints = {s["id"]: file_sha(CACHE / s["file"]) for s in SOURCES if s["id"] != 'kaikki-simple'}
    fingerprints["selectionVersion"] = "unicode-token-v3-retained-publication-history"
    if cache.exists() and (saved := json.loads(cache.read_text(encoding="utf-8"))).get("fingerprints") == fingerprints:
        contexts, collection = saved["contexts"], saved["collection"]
    else:
        contexts, collection = collect_contexts({w["key"] for w in words} | set(member_index))
        atomic_json(cache, {"fingerprints": fingerprints, "contexts": contexts, "collection": collection})
    exact_contexts = contexts
    contexts = add_dictionary_contexts(contexts, dictionary)
    contexts = add_editorial_contexts(contexts, dictionary)
    previous_path = OUT / 'entries.json'
    previous_entries = json.loads(previous_path.read_text(encoding='utf-8'))['entries'] if previous_path.exists() else []
    retain_published_contexts(contexts, dictionary, previous_entries)
    eligible = [w for w in words if contexts.get(w["key"])]
    ranked_keys = {w['key'] for w in words}
    eligible += [{'key': key, 'word': key, 'rank': {'sourceId': None, 'value': None, 'metric': None}} for key in sorted(member_index) if key not in ranked_keys and contexts.get(key)]
    previous_words = [normalized_word(e['word']) for e in previous_entries]
    selected = select_words(eligible, member_index, limit, previous_words)
    entries = []
    for item in selected:
        key = item["key"]
        reviewed = [c for c in contexts[key] if c['quality'] == 'context-reviewed']
        entries.append({"id": lexical_id(key), "kind": "word", "word": item["word"], "rank": item["rank"], "memberships": member_index.get(key, []), "cefr": None, "senses": dictionary.get(key, []), "contexts": contexts[key], "collocations": [p for c in reviewed for p in c.get('collocations', [])], "topics": [t for c in reviewed for t in c.get('topics', [])], "images": [], "quality": {"status": "context-reviewed" if reviewed else "imported-context", "reviewedContextCount": len(reviewed), "meaningAlignment": "reviewed-contexts-only" if reviewed else "not-reviewed", "americanEnglish": "reviewed-contexts-only" if reviewed else "not-reviewed", "register": "not-reviewed"}})
        editorial_sense = next((s for s in dictionary.get(key, []) if s['quality'] == 'editor-reviewed-for-this-context'), None)
        if editorial_sense:
            entries[-1]['lexicalType'] = editorial_sense['lexicalType']
            entries[-1]['displayHeadword'] = editorial_sense['displayHeadword']
    repairs_path = OUT / 'token-selection-repairs.json'
    repairs = None
    if repairs_path.exists():
        from complete_context_lexicon import apply_selection_repairs
        repairs = json.loads(repairs_path.read_text(encoding='utf-8'))
        entries = apply_selection_repairs({'entries': entries}, repairs)['entries']
    completion_path = OUT / 'editorial-completion.json'
    completion = None
    if completion_path.exists():
        from complete_context_lexicon import apply_overlay
        completion = json.loads(completion_path.read_text(encoding='utf-8'))
        entries = apply_overlay({'entries': entries}, completion)['entries']
        contexts.update({normalized_word(e['word']): e['contexts'] for e in entries})
    coverage = {"version": VERSION, "requestedWords": limit, "availableRankedCandidates": len(words), "availableWithLinkedContexts": len(exact_contexts), "availableWithAnyContext": len(eligible), "selectedWords": len(entries), "withDefinitions": sum(bool(e['senses']) for e in entries), "withRussianContexts": sum(any(c.get('ru') for c in e['contexts']) for e in entries), "contexts": sum(len(e['contexts']) for e in entries), "senses": sum(len(e['senses']) for e in entries), "editorReviewedWords": 0, "usReviewedWords": 0, "images": 0, "cefrRatedWords": 0, "core10000WithoutContext": [w['word'] for w in words[:10000] if w['key'] not in contexts], **collection}
    ranking = {w['key']: w for w in words}
    list_coverage = []
    for source_id in ['ngsl-1.2', 'ngsl-gr-1.0', 'ngsl-spoken-1.2', 'ngsl-supplement', 'nawl-1.2']:
        required = {word for word, memberships_ in member_index.items() if any(m['sourceId'] == source_id for m in memberships_)}
        missing = []
        for word in required:
            if word not in contexts:
                missing.append({'word': word, 'rank': ranking.get(word, {}).get('rank'), 'memberships': member_index[word], 'hasDefinition': bool(dictionary.get(word)), 'hasDictionaryExample': any(s.get('examples') for s in dictionary.get(word, [])), 'reason': 'no selected direct EN/RU exact-token pair or suitable dictionary example; all required list headwords scanned'})
        missing.sort(key=lambda x: (x['rank']['value'] if x['rank'] else 999999, x['word']))
        selected_keys = {w['key'] for w in selected}
        list_coverage.append({'sourceId': source_id, 'requiredWords': len(required), 'withContext': sum(w in contexts for w in required), 'inBankWithContext': len(required & selected_keys), 'withRussianContext': sum(any(c.get('ru') for c in contexts.get(w, [])) for w in required), 'missingCount': len(missing), 'missing': missing})
    coverage['listCoverage'] = [{k: v for k, v in group.items() if k != 'missing'} for group in list_coverage]
    gap_document = {'version': VERSION, 'stage': 'exact-token EN/RU selection plus source-linked dictionary examples and explicitly reviewed editorial additions; not a claim that no usable example exists anywhere', 'criteria': {'englishWords': [7, 28], 'maxEnglishCharacters': 240, 'match': 'case-insensitive complete token; apostrophe/hyphen are inside tokens; no stemming', 'translation': 'direct Tatoeba English→Russian link', 'selection': '4 candidates per word scored by length, then source ID; at most2 returned; formal filters only', 'dictionaryFallback': '5–35 words, no external ref/quotation, sense-assigned example with lemma or explicit same-entry inflected form', 'editorial': 'original scenario with explicit selected meaning, RU translation, explanation and production task; approved files only', 'variety': 'imported contexts not automatically verified as American'}, 'lists': list_coverage, 'rankedFirst10000Gaps': [w for w in words[:10000] if w['key'] not in contexts], 'first80RankedGaps': [w for w in words[:10000] if w['key'] not in contexts][:80]}
    coverage['editorReviewedWords'] = sum(e['quality']['reviewedContextCount'] > 0 for e in entries)
    coverage['usReviewedWords'] = sum(any(c['quality'] == 'context-reviewed' and c.get('variety') == 'en-US-compatible' for c in e['contexts']) for e in entries)
    coverage['reviewedContexts'] = sum(e['quality']['reviewedContextCount'] for e in entries)
    coverage['aiReviewedContexts'] = sum(c.get('quality') == 'ai-context-reviewed' for e in entries for c in e['contexts'])
    coverage['aiReviewedWords'] = sum(any(c.get('quality') == 'ai-context-reviewed' for c in e['contexts']) for e in entries)
    coverage['contextsWithoutRussian'] = sum(not c.get('ru') for e in entries for c in e['contexts'])
    coverage['studyContextsWithoutRussian'] = sum(not c.get('ru') and not c.get('excludedFromStudy') for e in entries for c in e['contexts'])
    coverage['contextsWithoutSelectedMeaning'] = sum(not c.get('senseId') for e in entries for c in e['contexts'])
    coverage['studyContextsWithoutSelectedMeaning'] = sum(not c.get('senseId') and not c.get('excludedFromStudy') for e in entries for c in e['contexts'])
    coverage['entriesWithTranslatedSelectedMeaning'] = sum(any(c.get('ru') and c.get('senseId') and not c.get('excludedFromStudy') for c in e['contexts']) for e in entries)
    coverage['entriesWithoutStudyContext'] = sum(not any(not c.get('excludedFromStudy') for c in e['contexts']) for e in entries)
    coverage['excludedSourceContexts'] = sum(bool(c.get('excludedFromStudy')) for e in entries for c in e['contexts'])
    coverage['studyContexts'] = sum(not c.get('excludedFromStudy') for e in entries for c in e['contexts'])
    coverage['withSourceContextIssues'] = sum(any(c.get('sourceIssues') for c in e['contexts']) for e in entries)
    coverage['aiReviewScope'] = 'Only named contexts in editorial-completion.json; separate AI drafting and semantic review, not full human review'
    coverage['withSourceIssues'] = sum(any(s.get('sourceIssues') for s in e['senses']) for e in entries)
    coverage['countingUnit'] = 'distinct source headword records; spelling variants and combining forms are explicitly marked, not counted as independently mastered lemmas'
    coverage['combiningForms'] = sum(e.get('lexicalType') == 'combining-form' for e in entries)
    coverage['markedSpellingVariants'] = sum(e.get('lexicalType') == 'spelling-variant' for e in entries)
    coverage['scope'] = 'available NGSL/GR/Spoken/supplement/NAWL words first, then available NGSL31k candidates by raw frequency; missing ranked/core words remain reported separately'
    coverage['published'] = not pilot
    pilot_entries = [e for e in entries if e['quality']['reviewedContextCount']][:20]
    # Pilot presents only reviewed contexts/senses; full bank retains all source references.
    if pilot:
        pilot_entries = json.loads(json.dumps(pilot_entries))
        for entry in pilot_entries:
            entry['contexts'] = [c for c in entry['contexts'] if c['quality'] == 'context-reviewed']
            ids = {c['senseId'] for c in entry['contexts']}
            entry['senses'] = [s for s in entry['senses'] if s['id'] in ids]
    receipts = [json.loads((CACHE / (s['file'] + '.receipt.json')).read_text(encoding='utf-8')) for s in SOURCES]
    source_records = [*receipts, *[{**data['source'], 'file': path.name, 'sha256': file_sha(path)} for path, data in editorial_documents()]]
    if completion:
        source_records.append({**completion['source'], 'file': completion_path.name, 'sha256': file_sha(completion_path)})
    if repairs:
        source_records.append({**repairs['source'], 'file': repairs_path.name, 'sha256': file_sha(repairs_path)})
    document = {"version": VERSION, "targetVariety": "en-US", "spanEncoding": "utf-16", "entries": pilot_entries if pilot else entries}
    if repairs:
        document['unicodeTokenSelectionValidated'] = True
    validate_entries(document, source_records)
    atomic_json(OUT / ('pilot-sources.json' if pilot else 'sources.json'), {'version': VERSION, 'sources': source_records})
    atomic_json(OUT / ("pilot.json" if pilot else "entries.json"), document)
    coverage['outputFile'] = 'pilot.json' if pilot else 'entries.json'
    coverage['outputSHA256'] = file_sha(OUT / coverage['outputFile'])
    coverage['publishedWords'] = 0 if pilot else len(entries)
    # Publish coverage only after the validated dataset exists; consumers can
    # verify the output digest when files change between separate reads.
    atomic_json(OUT / ('pilot-required-gaps.json' if pilot else 'required-gaps.json'), gap_document)
    atomic_json(OUT / ('pilot-coverage.json' if pilot else 'coverage.json'), coverage)
    print(json.dumps({k: v for k, v in coverage.items() if not isinstance(v, list)}, ensure_ascii=True), flush=True)
    return entries, coverage


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--download", action="store_true")
    parser.add_argument("--refresh-sources", action="store_true")
    parser.add_argument("--build", action="store_true")
    parser.add_argument("--pilot", action="store_true")
    parser.add_argument("--limit", type=int, default=10000)
    args = parser.parse_args()
    if args.download:
        download_sources(args.refresh_sources)
    if args.build or args.pilot:
        build(args.limit, args.pilot)
